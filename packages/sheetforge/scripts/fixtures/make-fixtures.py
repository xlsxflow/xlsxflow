"""Builds reader test fixtures + expected values (test/fixtures/*.xlsx, *.expected.json).

  python scripts/fixtures/make-fixtures.py            # openpyxl + hand-crafted fixtures, then oracles
  python scripts/fixtures/make-fixtures.py --oracle   # only (re)write oracles for every .xlsx lacking a crafted one

Oracle: openpyxl reads each file; values are normalized to SheetForge's conventions:
dates/times -> ISO-8601 UTC strings (Excel serial semantics, see iso()), errors -> '#DIV/0!' etc.,
trailing empty cells trimmed, empty rows dropped. Crafted fixtures carry hand-written expectations.
Requires: pip install openpyxl
"""
import datetime as dt
import json
import sys
import zipfile
from pathlib import Path

import openpyxl
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.utils.datetime import CALENDAR_MAC_1904, CALENDAR_WINDOWS_1900, to_excel

FIX = Path(__file__).resolve().parents[2] / 'test' / 'fixtures'


# ---------- normalization shared by the oracle and crafted expectations ----------

def iso(serial: float, is1904: bool) -> str:
    """Excel serial -> ISO string, matching SheetForge: 1900 leap-year bug honoured, ms precision."""
    days = int(serial // 1)
    ms = round((serial - days) * 86400000)
    if is1904:
        base = dt.datetime(1904, 1, 1)
    else:
        base = dt.datetime(1899, 12, 30)
        if 1 <= days <= 60:
            days += 1
    t = base + dt.timedelta(days=days, milliseconds=ms)
    return t.strftime('%Y-%m-%dT%H:%M:%S.') + f'{t.microsecond // 1000:03d}Z'


def norm(v, is1904):
    epoch = CALENDAR_MAC_1904 if is1904 else CALENDAR_WINDOWS_1900
    if isinstance(v, dt.timedelta):
        return iso(v.total_seconds() / 86400, is1904)
    if isinstance(v, (dt.datetime, dt.date, dt.time)):
        return iso(to_excel(v, epoch), is1904)
    if isinstance(v, float) and v.is_integer() and abs(v) < 2**53:
        return int(v)
    return v


def trim(cells):
    while cells and cells[-1] is None:
        cells.pop()
    return cells


def oracle(path: Path) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True)
    is1904 = wb.epoch == CALENDAR_MAC_1904
    sheets = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows():
            cells = trim([norm(c.value, is1904) for c in row])
            if cells:
                rows.append([row[0].row, cells])
        sheets.append({'name': ws.title, 'rows': rows})
    return {'source': 'openpyxl oracle', 'sheets': sheets}


def write_expected(path: Path, expected: dict):
    path.with_suffix('.expected.json').write_text(json.dumps(expected, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')


# ---------- openpyxl-written fixtures ----------

def openpyxl_book(is1904: bool) -> openpyxl.Workbook:
    wb = openpyxl.Workbook()
    if is1904:
        wb.epoch = CALENDAR_MAC_1904
    t = wb.active
    t.title = 'Types'
    t.append(['text', 42, -3.5, 1e20, 1e-10, 0.1 + 0.2])
    t.append([True, False, None, None, '#N/A literal text'])
    t['A4'] = 'row 4 after a gap'
    t['Z4'] = 'far column'
    t['A5'] = 'emoji \U0001F600 arabic \u0645\u0631\u062d\u0628\u0627 cjk \u6f22\u5b57'
    t['B5'] = '  padded  '
    t['C5'] = 'x & y < z > "q" \'s\''
    t['D5'] = 'line1\nline2'
    t['E5'] = '00123'
    t['F5'] = 'L' * 32767
    t['G5'] = 'text'

    d = wb.create_sheet('Dates')
    cells = [
        ('A1', dt.datetime(2024, 2, 29), 'yyyy-mm-dd'),
        ('B1', dt.datetime(2024, 2, 29, 18, 0), 'dd/mm/yyyy hh:mm'),
        ('C1', dt.time(12, 0), 'hh:mm:ss'),
        ('D1', dt.datetime(2023, 3, 15), '[$-409]mmmm d, yyyy;@'),
        ('A2', 0.25, '0.00%'),
        ('B2', 1234.5, '#,##0.00 "USD"'),
        ('C2', dt.timedelta(hours=36), '[h]:mm'),
        ('D2', 7, '0 "days"'),
        ('E2', 123.456, '[Red]0.0;[Blue]-0.0'),
        ('F2', 5, '"Year" 0'),
    ]
    for ref, val, fmt in cells:
        d[ref] = val
        d[ref].number_format = fmt

    r = wb.create_sheet('Rich & Layout')
    r['A1'] = CellRichText([TextBlock(InlineFont(b=True), 'bold'), ' then plain'])
    r['A2'] = 'merged'
    r.merge_cells('A2:C3')
    r['A5'] = 'hidden row'
    r.row_dimensions[5].hidden = True
    r['D1'] = 'hidden column'
    r.column_dimensions['D'].hidden = True
    r.freeze_panes = 'A2'

    b = wb.create_sheet('Bulk')
    for i in range(5000):   # enough to span many decompression chunks
        b.append([i, f'name {i % 500}', i * 0.5, i % 2 == 0, f'unique-{i}'])

    wb.create_sheet('Empty')
    wb.move_sheet('Bulk', offset=-3)   # tab order != part order
    wb.active = 0
    return wb


# ---------- hand-crafted fixtures (edge cases no library above produces) ----------

MAIN = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
REL = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
PKG = 'http://schemas.openxmlformats.org/package/2006/relationships'


def craft(name, parts: dict, expected_sheets, stored=False):
    path = FIX / name
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_STORED if stored else zipfile.ZIP_DEFLATED) as z:
        for part, text in parts.items():
            z.writestr(part, text.encode('utf-8') if isinstance(text, str) else text)
    write_expected(path, {'source': 'hand-written', 'sheets': expected_sheets})


def minimal(sheet_xml, *, shared=None, styles=None, workbook_extra='', sheet_name='S', main=MAIN, rel=REL):
    parts = {
        '[Content_Types].xml': '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>',
        '_rels/.rels': f'<Relationships xmlns="{PKG}"><Relationship Id="rId1" Type="{rel}/officeDocument" Target="xl/workbook.xml"/></Relationships>',
        'xl/workbook.xml': f'<workbook xmlns="{main}" xmlns:r="{rel}">{workbook_extra}<sheets><sheet name="{sheet_name}" sheetId="1" r:id="rId1"/></sheets></workbook>',
        'xl/_rels/workbook.xml.rels': f'<Relationships xmlns="{PKG}"><Relationship Id="rId1" Type="{rel}/worksheet" Target="worksheets/sheet1.xml"/>'
                                      f'<Relationship Id="rId2" Type="{rel}/sharedStrings" Target="sharedStrings.xml"/>'
                                      f'<Relationship Id="rId3" Type="{rel}/styles" Target="styles.xml"/></Relationships>',
        'xl/worksheets/sheet1.xml': sheet_xml,
    }
    if shared is not None:
        parts['xl/sharedStrings.xml'] = f'<sst xmlns="{main}">{shared}</sst>'
    if styles is not None:
        parts['xl/styles.xml'] = f'<styleSheet xmlns="{main}">{styles}</styleSheet>'
    return parts


def ws(body, main=MAIN):
    return f'<?xml version="1.0" encoding="UTF-8"?><worksheet xmlns="{main}"><sheetData>{body}</sheetData></worksheet>'


def crafted():
    # Strict OOXML namespaces (Excel "Strict Open XML Spreadsheet")
    strict_main = 'http://purl.oclc.org/ooxml/spreadsheetml/main'
    strict_rel = 'http://purl.oclc.org/ooxml/officeDocument/relationships'
    craft('crafted-strict.xlsx', minimal(
        ws('<row r="1"><c r="A1" t="s"><v>0</v></c><c r="B1"><v>1.5</v></c></row>', main=strict_main),
        shared='<si><t>strict</t></si>', main=strict_main, rel=strict_rel),
        [{'name': 'S', 'rows': [[1, ['strict', 1.5]]]}])

    # Every element namespace-prefixed (e.g. .NET Open XML SDK output)
    craft('crafted-prefixed.xlsx', {
        **minimal(''),
        'xl/workbook.xml': f'<x:workbook xmlns:x="{MAIN}" xmlns:r="{REL}"><x:sheets><x:sheet name="P" sheetId="1" r:id="rId1"/></x:sheets></x:workbook>',
        'xl/worksheets/sheet1.xml': f'<x:worksheet xmlns:x="{MAIN}"><x:sheetData><x:row r="1"><x:c r="A1" t="inlineStr"><x:is><x:t>pre</x:t></x:is></x:c>'
                                    f'<x:c r="B1"><x:v>2</x:v></x:c></x:row></x:sheetData></x:worksheet>',
    }, [{'name': 'P', 'rows': [[1, ['pre', 2]]]}])

    # Non-standard part names: workbook, sheets, shared strings found only via relationships
    craft('crafted-custom-paths.xlsx', {
        '[Content_Types].xml': '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>',
        '_rels/.rels': f'<Relationships xmlns="{PKG}"><Relationship Id="r" Type="{REL}/officeDocument" Target="/book/main.xml"/></Relationships>',
        'book/main.xml': f'<workbook xmlns="{MAIN}" xmlns:r="{REL}"><sheets><sheet name="Only" sheetId="7" r:id="a"/></sheets></workbook>',
        'book/_rels/main.xml.rels': f'<Relationships xmlns="{PKG}"><Relationship Id="a" Type="{REL}/worksheet" Target="data/first.xml"/>'
                                    f'<Relationship Id="b" Type="{REL}/sharedStrings" Target="strings.xml"/></Relationships>',
        'book/data/first.xml': ws('<row r="1"><c r="A1" t="s"><v>0</v></c></row>'),
        'book/strings.xml': f'<sst xmlns="{MAIN}"><si><t>found via rels</t></si></sst>',
    }, [{'name': 'Only', 'rows': [[1, ['found via rels']]]}])

    # Shared strings: rich runs, phonetic runs, _xHHHH_ escapes, entities, preserved whitespace
    craft('crafted-strings.xlsx', minimal(
        ws(''.join(f'<row r="{i + 1}"><c r="A{i + 1}" t="s"><v>{i}</v></c></row>' for i in range(7))
           + '<row r="8"><c r="A8" t="inlineStr"><is><r><t>in</t></r><r><t xml:space="preserve"> line</t></r></is></c>'
             '<c r="B8" t="inlineStr"><is><t>a_x000D_b</t></is></c></row>'),
        shared='<si><r><rPr><b/></rPr><t>bold</t></r><r><t xml:space="preserve"> plain</t></r></si>'
               '<si><t>\u6f22\u5b57</t><rPh sb="0" eb="2"><t>\u30ab\u30f3\u30b8</t></rPh></si>'
               '<si><t>tab_x0009_cr_x000D_</t></si>'
               '<si><t>_x005F_x0041_ stays literal</t></si>'
               '<si><t>&lt;&amp;&gt; &#x1F600; &#65;</t></si>'
               '<si><t xml:space="preserve">  spaced  </t></si>'
               '<si><t/></si>'),
        [{'name': 'S', 'rows': [[1, ['bold plain']], [2, ['\u6f22\u5b57']], [3, ['tab\tcr\r']],
                                [4, ['_x0041_ stays literal']], [5, ['<&> \U0001F600 A']], [6, ['  spaced  ']],
                                [8, ['in line', 'a\rb']]]}])
    # Row 7 is an empty shared string: an empty-string cell. Readers differ (openpyxl: None); we keep ''.
    exp = json.loads((FIX / 'crafted-strings.expected.json').read_text(encoding='utf-8'))
    exp['sheets'][0]['rows'].insert(6, [7, ['']])
    write_expected(FIX / 'crafted-strings.xlsx', exp)

    # Cells/rows without r attributes, explicit types, odd numerics, shared formulas, AlternateContent
    craft('crafted-cells.xlsx', minimal(
        f'<worksheet xmlns="{MAIN}" xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"><sheetData>'
        '<row><c><v>1</v></c><c t="n"><v>1.5E+20</v></c><c><v>-0</v></c><c t="str"><v>formula text</v></c></row>'
        '<row><c t="b"><v>1</v></c><c t="e"><v>#REF!</v></c><c t="d"><v>2024-01-02T03:04:05Z</v></c><c/><c t="s"/><c><v></v></c></row>'
        '<row r="5"><c r="C5"><f t="shared" ref="C5:C6" si="0">A1*2</f><v>2</v></c></row>'
        '<row r="6"><c r="C6"><f t="shared" si="0"/><v>4</v></c></row>'
        '<mc:AlternateContent><mc:Choice Requires="x14"><row r="99"><c r="A99"><v>999</v></c></row></mc:Choice></mc:AlternateContent>'
        '<row r="7"><c r="XFD7"><v>7</v></c></row>'
        '</sheetData></worksheet>'),
        [{'name': 'S', 'rows': [[1, [1, 1.5e20, 0, 'formula text']],
                                [2, [True, '#REF!', '2024-01-02T03:04:05Z']],
                                [5, [None, None, 2]], [6, [None, None, 4]],
                                [7, [None] * 16383 + [7]]]}])

    # Dates: builtin + custom formats, 1900 leap-year bug, quoted/bracketed literals that are NOT dates
    fmts = ('<numFmts count="5"><numFmt numFmtId="164" formatCode="0 &quot;days&quot;"/>'
            '<numFmt numFmtId="165" formatCode="[Red]0.0"/><numFmt numFmtId="166" formatCode="yyyy\\-mm\\-dd"/>'
            '<numFmt numFmtId="167" formatCode="[h]:mm"/><numFmt numFmtId="168" formatCode="\\d0"/></numFmts>')
    xfs = ''.join(f'<xf numFmtId="{n}"/>' for n in (0, 14, 22, 164, 165, 166, 167, 168, 10, 47))
    cells = [(1, 45000), (1, 1), (1, 59), (1, 60), (1, 61), (2, 45351.75), (1, 0.5),
             (3, 7), (4, 5), (5, 45000), (6, 1.5), (7, 3), (8, 0.25), (9, 0.5)]
    body = ''.join(f'<row r="{i + 1}"><c r="A{i + 1}" s="{s}"><v>{v}</v></c></row>' for i, (s, v) in enumerate(cells))
    date_styles = {1, 2, 5, 6, 9}
    rows = [[i + 1, [iso(v, False) if s in date_styles else v]] for i, (s, v) in enumerate(cells)]
    craft('crafted-dates.xlsx', minimal(ws(body), styles=f'{fmts}<cellXfs>{xfs}</cellXfs>'),
          [{'name': 'S', 'rows': rows}])
    craft('crafted-dates-1904.xlsx', minimal(ws(body), styles=f'{fmts}<cellXfs>{xfs}</cellXfs>',
                                             workbook_extra='<workbookPr date1904="1"/>'),
          [{'name': 'S', 'rows': [[i + 1, [iso(v, True) if s in date_styles else v]] for i, (s, v) in enumerate(cells)]}])

    # Stored (uncompressed) entries, BOM-prefixed XML, escaped sheet name, empty self-closing sheetData
    parts = minimal('\ufeff' + ws('<row r="1"><c r="A1"><v>1</v></c></row>'), sheet_name='A&amp;B &quot;q&quot;')
    craft('crafted-stored-bom.xlsx', parts, [{'name': 'A&B "q"', 'rows': [[1, [1]]]}], stored=True)
    craft('crafted-empty.xlsx', minimal(f'<worksheet xmlns="{MAIN}"><sheetData/></worksheet>'),
          [{'name': 'S', 'rows': []}])


def main():
    FIX.mkdir(parents=True, exist_ok=True)
    if '--oracle' not in sys.argv:
        openpyxl_book(False).save(FIX / 'openpyxl-basic.xlsx')
        openpyxl_book(True).save(FIX / 'openpyxl-1904.xlsx')
        crafted()
    for f in sorted(FIX.glob('*.xlsx')):
        exp = f.with_suffix('.expected.json')
        # Hand-written / source-data expectations are never overwritten by the oracle
        if exp.exists() and json.loads(exp.read_text(encoding='utf-8')).get('source') != 'openpyxl oracle':
            continue
        try:
            write_expected(f, oracle(f))
            print('oracle', f.name)
        except Exception as e:
            print(f'NO ORACLE for {f.name}: openpyxl cannot read it ({type(e).__name__}); supply expectations another way')


if __name__ == '__main__':
    main()
