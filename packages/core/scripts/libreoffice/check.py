"""Checks edited workbooks against Excel's file schema and in LibreOffice.

Schema check: Microsoft's Open XML SDK (scripts/openxml/validate.cs) checks each file built by make.ts
against the rules Excel is strict about. A file edited from an openpyxl workbook only fails on errors
its source does not already have.

LibreOffice check: LibreOffice opens each file, recalculates every formula itself (the files' cached
values are ignored) and saves a copy. It passes when every expected value matches and the copy keeps
the same merges, rules, validations, tables, notes, pictures, charts, filters and protection.

Needs Python with openpyxl and Node (npx tsx), plus the .NET 10 SDK for the schema check and
LibreOffice for the other; a check whose tool is missing is skipped. Set SOFFICE when soffice is not
on PATH or in its default place. Run from anywhere:

    python packages/core/scripts/libreoffice/check.py
"""
import json
import os
import pathlib
import shutil
import subprocess
import sys

import openpyxl
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.filters import FilterColumn, Filters
from openpyxl.worksheet.pagebreak import Break

HERE = pathlib.Path(__file__).resolve().parent
OUT = HERE / 'out'


def find_soffice():
    for candidate in (os.environ.get('SOFFICE'), shutil.which('soffice'),
                      r'C:\Program Files\LibreOffice\program\soffice.com',
                      '/Applications/LibreOffice.app/Contents/MacOS/soffice'):
        if candidate and pathlib.Path(candidate).exists():
            return candidate
    return None


def openpyxl_sources():
    """Workbooks made by another library, with a chart, filters, page breaks and print titles."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Data'
    ws.append(['Name', 'Qty'])
    for i in range(1, 11):
        ws.append([f'item {i}', i * 10])
    ws['D1'] = '=SUM(B2:B11)'
    ws['A6'].comment = Comment('row 6 note', 'Ana')
    ws.merge_cells('D5:E6')
    ws.conditional_formatting.add('B2:B11', CellIsRule(operator='greaterThan', formula=['50'], fill=PatternFill('solid', fgColor='FFC7CE')))
    dv = DataValidation(type='whole', operator='between', formula1='0', formula2='1000')
    dv.add('B2:B11')
    ws.add_data_validation(dv)
    ws.auto_filter.ref = 'A1:B11'
    ws.freeze_panes = 'A2'
    ws.print_title_rows = '1:1'
    ws.row_breaks.append(Break(id=8))
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=11), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=11))
    ws.add_chart(chart, 'G8')
    other = wb.create_sheet('Other')
    other['A1'] = '=Data!B8'
    other['A2'] = '=SUM(Data!B2:B11)'
    wb.save(OUT / 'openpyxl-rows.xlsx')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Data'
    ws.append(['Name', 'Region', 'Qty', 'Price'])
    for i in range(1, 9):
        ws.append([f'item {i}', 'N' if i % 2 else 'S', i * 10, i * 1.5])
    ws['F1'] = '=SUM(C2:C9)'
    ws['F2'] = '=C2*D2'
    ws['C3'].comment = Comment('qty note', 'Ana')
    ws.merge_cells('B11:D11')
    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['C'].width = 12
    ws.conditional_formatting.add('C2:C9', CellIsRule(operator='greaterThan', formula=['40'], fill=PatternFill('solid', fgColor='FFC7CE')))
    dv = DataValidation(type='decimal', operator='greaterThan', formula1='0')
    dv.add('D2:D9')
    ws.add_data_validation(dv)
    ws.auto_filter.ref = 'A1:D9'
    ws.auto_filter.filterColumn.append(FilterColumn(colId=1, filters=Filters(filter=['N'])))
    ws.col_breaks.append(Break(id=3))
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=3, min_row=1, max_row=9), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=9))
    ws.add_chart(chart, 'H2')
    ws.print_title_cols = 'A:A'
    other = wb.create_sheet('Other')
    other['A1'] = '=Data!D5'
    other['A2'] = '=SUM(Data!C2:D9)'
    wb.save(OUT / 'openpyxl-cols.xlsx')


def features(path):
    wb = openpyxl.load_workbook(path)
    return {ws.title: {
        'merges': sorted(str(r) for r in ws.merged_cells.ranges),
        'conditional formats': sorted(str(r.sqref) for r in ws.conditional_formatting),
        'validations': sorted(str(v.sqref) for v in ws.data_validations.dataValidation),
        'tables': sorted((t.name, t.ref) for t in ws.tables.values()),
        'notes': sorted(c.coordinate for row in ws.iter_rows() for c in row if c.comment),
        'pictures': len(ws._images),
        'charts': len(ws._charts),
        'filter': ws.auto_filter.ref,
        'protected': ws.protection.sheet,
    } for ws in wb}


def same_value(got, want):
    if isinstance(want, (int, float)) and not isinstance(want, bool) and isinstance(got, (int, float)):
        return abs(got - want) < 1e-9
    # A reference to deleted cells on another sheet is stored as Data!#REF!, as Excel stores it;
    # LibreOffice 24.2 evaluates that form to #NAME? instead of #REF!
    if want == '#REF!':
        return got in ('#REF!', '#NAME?')
    return got == want


def schema_errors(paths):
    """{file name: [error]} from the Open XML SDK"""
    run = subprocess.run(['dotnet', 'run', 'validate.cs', '--', *map(str, paths)],
                         cwd=HERE.parent / 'openxml', capture_output=True, text=True)
    errors, current = {}, None
    for line in run.stdout.splitlines():
        if line.startswith(('ok   ', 'FAIL ')):
            current = line[5:].split(':')[0]
            errors[current] = [line[5 + len(current) + 2:]] if ':' in line[5:] else []
        elif line.startswith('     ') and current:
            errors[current].append(line.strip())
    if not errors:
        sys.exit('Schema check could not run:\n' + run.stdout + run.stderr)
    return errors


def schema_check(expected):
    errors = schema_errors([OUT / name for name in expected] + [OUT / 'openpyxl-rows.xlsx', OUT / 'openpyxl-cols.xlsx'])
    failures = 0
    for name in expected:
        source = 'openpyxl-rows.xlsx' if name.startswith('openpyxl-rows') else 'openpyxl-cols.xlsx' if name.startswith('openpyxl-cols') else None
        new = [e for e in errors[name] if not source or e not in errors[source]]
        print(('ok   ' if not new else 'FAIL ') + name + ''.join('\n     ' + e for e in new[:10]))
        failures += bool(new)
    print(f'{len(expected) - failures} of {len(expected)} files passed the schema check')
    return failures


def libreoffice_check(soffice, expected):
    # A private profile that recalculates every formula when an xlsx file is loaded
    profile = OUT / 'profile'
    (profile / 'user').mkdir(parents=True)
    (profile / 'user' / 'registrymodifications.xcu').write_text(
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<oor:items xmlns:oor="http://openoffice.org/2001/registry" xmlns:xs="http://www.w3.org/2001/XMLSchema">'
        '<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="OOXMLRecalcMode" oor:op="fuse">'
        '<value>0</value></prop></item></oor:items>\n')

    saved = OUT / 'libreoffice'
    failures = 0
    for name, cells in expected.items():
        subprocess.run([soffice, '--headless', '--norestore', f'-env:UserInstallation={profile.as_uri()}',
                        '--convert-to', 'xlsx', '--outdir', str(saved), str(OUT / name)], capture_output=True, timeout=300)
        copy = saved / name
        if not copy.exists():
            print(f'FAIL {name}: LibreOffice could not open it')
            failures += 1
            continue
        wb = openpyxl.load_workbook(copy, data_only=True)
        problems = []
        for ref, want in cells.items():
            sheet, cell = ref.split('!')
            got = wb[sheet][cell].value
            if not same_value(got, want):
                problems.append(f'{ref} is {got!r}, expected {want!r}')
        ours, theirs = features(OUT / name), features(copy)
        for sheet, found in ours.items():
            for feature, value in found.items():
                other = theirs.get(sheet, {}).get(feature)
                if other != value:
                    problems.append(f'{sheet} {feature}: {value!r} became {other!r}')
        print(('ok   ' if not problems else 'FAIL ') + name + ('' if not problems else ': ' + '; '.join(problems)))
        failures += bool(problems)
    print(f'{len(expected) - failures} of {len(expected)} files passed in LibreOffice')
    return failures


def main():
    shutil.rmtree(OUT, ignore_errors=True)
    OUT.mkdir()
    openpyxl_sources()
    subprocess.run('npx -y tsx@4.23.15 scripts/libreoffice/make.ts scripts/libreoffice/out', shell=True, check=True, cwd=HERE.parents[1])
    expected = json.loads((OUT / 'expected.json').read_text())

    failures, ran = 0, 0
    if shutil.which('dotnet'):
        failures += schema_check(expected)
        ran += 1
    else:
        print('Schema check skipped: the .NET 10 SDK (dotnet) was not found')
    soffice = find_soffice()
    if soffice:
        failures += libreoffice_check(soffice, expected)
        ran += 1
    else:
        print('LibreOffice check skipped: LibreOffice was not found (install it or set SOFFICE)')
    if not ran:
        sys.exit('No check ran')
    if os.environ.get('CI') and ran < 2:
        sys.exit('CI runs both checks')
    sys.exit(1 if failures else 0)


if __name__ == '__main__':
    main()
